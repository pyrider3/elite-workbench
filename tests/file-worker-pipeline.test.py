import unittest,tempfile,pathlib,importlib.util,json,zipfile
spec=importlib.util.spec_from_file_location('p',pathlib.Path(__file__).parents[1]/'scripts/file-pipeline.py');p=importlib.util.module_from_spec(spec);spec.loader.exec_module(p)
class Tests(unittest.TestCase):
 def test_resume_after_render_failure_and_corruption(self):
  with tempfile.TemporaryDirectory() as t:
   root=pathlib.Path(t);f=root/'real.md';f.write_text('# 原标题\n保留2026年10月23日和电话123456。');calls=[]
   def ask(prompt):calls.append(1);return json.dumps({'edits':[{'id':'0:text','text':'# 新标题\n保留2026年10月23日和电话123456。'}],'creates':[]})
   pipe=p.Pipeline(root/'cache',[f],'改标题');out=pipe.run(ask);self.assertEqual(len(calls),1);original=f.read_text()
   out[0].write_text('corrupt');out=pipe.run(lambda _:self.fail('must reuse plan'));self.assertIn('新标题',out[0].read_text());self.assertEqual(f.read_text(),original)
   self.assertEqual(pipe.run(lambda _:self.fail('must reuse rendered'))[0].read_text(),out[0].read_text())
 def test_invalid_ids_and_injection(self):
  docs=[{'units':[{'id':'0:1','text':'x'},{'id':'0:2','text':'=SUM(A1)','readonly':True}]}]
  for plan in [{'edits':[{'id':'other','text':'x'}],'creates':[]},{'edits':[{'id':'0:2','text':'x'}],'creates':[]},{'edits':[],'creates':[{'name':'../bad.docx','text':'bad'}]}]:
   with self.assertRaises(p.PipelineError):p.parse_plan(json.dumps(plan),docs)
 def test_office_run_preservation(self):
  from docx import Document
  with tempfile.TemporaryDirectory() as t:
   root=pathlib.Path(t);f=root/'source.docx';d=Document();para=d.add_paragraph();para.add_run('重要：').bold=True;para.add_run('原文2026');d.add_table(rows=1,cols=1).cell(0,0).text='保留表格';d.save(f);before=p.sha(f)
   docs=p.extract([f]);plan=p.parse_plan(json.dumps({'edits':[{'id':'0:0','text':'重要：新文2026'}],'creates':[]}),docs);out=p.render([f],docs,plan,root/'out')[0];r=Document(out)
   self.assertEqual(r.paragraphs[0].text,'重要：新文2026');self.assertTrue(r.paragraphs[0].runs[0].bold);self.assertEqual(r.tables[0].cell(0,0).text,'保留表格');self.assertEqual(p.sha(f),before)
 def test_spreadsheet_formula_and_styles_preserved(self):
  from openpyxl import Workbook,load_workbook
  with tempfile.TemporaryDirectory() as t:
   root=pathlib.Path(t);f=root/'source.xlsx';w=Workbook();ws=w.active;ws['A1']='原文';ws['A2']=123;ws['B2']='=A2*2';ws['A1'].number_format='@';w.save(f)
   docs=p.extract([f]);out=p.render([f],docs,{'edits':[{'id':'0:0','text':'修改'}],'creates':[]},root/'out')[0];r=load_workbook(out);self.assertEqual(r.active['B2'].value,'=A2*2');self.assertEqual(r.active['A1'].value,'修改');self.assertEqual(r.active['A1'].number_format,'@');r.close()
 def test_plan_saved_before_render(self):
  from unittest.mock import patch
  with tempfile.TemporaryDirectory() as t:
   root=pathlib.Path(t);f=root/'a.md';f.write_text('old');pipe=p.Pipeline(root/'cache',[f],'new')
   with patch.object(p,'render',side_effect=RuntimeError('power loss')):
    with self.assertRaises(RuntimeError):pipe.run(lambda _: '{"edits":[{"id":"0:text","text":"new"}],"creates":[]}')
   out=pipe.run(lambda _:self.fail('AI called twice'));self.assertEqual(out[0].read_text(),'new')
 def test_bounded_schema_repair(self):
  with tempfile.TemporaryDirectory() as t:
   root=pathlib.Path(t);f=root/'a.md';f.write_text('old');answers=iter(['not json','{"edits":[{"id":"0:text","text":"new"}],"creates":[]}']);calls=[]
   def ask(_):calls.append(1);return next(answers)
   result=p.Pipeline(root/'cache',[f],'new').run(ask);self.assertEqual(len(calls),2);self.assertEqual(result[0].read_text(),'new')
 def test_direct_routing(self):
  self.assertEqual(p.direct_kind([pathlib.Path('a.docx')],'转为pdf'),'pdf');self.assertEqual(p.direct_kind([pathlib.Path('a.xlsx')],'把这个做成张精美的展示的图片'),'poster');self.assertIsNone(p.direct_kind([pathlib.Path('a.docx')],'修改日期后转为pdf'))
if __name__=='__main__':unittest.main()
